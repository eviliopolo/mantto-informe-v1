from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from typing import Dict, List, Any, Optional
from datetime import datetime
import calendar
import json
import config
import os
from urllib.parse import urlparse
# Nota: get_glpi_extractor ahora es async, se debe llamar cuando se necesite

from docxtpl import DocxTemplate
from docx import Document
from pathlib import Path
from docxtpl import InlineImage
from docx.shared import Mm
import base64
from PIL import Image
from io import BytesIO
import logging
import re

# Importar utilidades
from src.utils.fecha_utils import formatear_fecha_simple, formatear_fechas_en_tabla
from src.utils.tabla_utils import (
    set_cell_shading, centrar_celda_vertical, enable_autofit,
    detectar_columnas_desde_datos, preparar_datos_tabla, count_rows,
    reemplazar_placeholder_con_tabla
)
from src.utils.imagen_utils import base64_to_inline_image, procesar_imagen
from src.utils.documento_utils import convertir_url_sharepoint_a_ruta_relativa

logger = logging.getLogger(__name__)


class GeneradorSeccion2:
    @property
    def nombre_seccion(self) -> str:
        return "2. INFORME DE MESA DE SERVICIO"
    
    @property
    def template_file(self) -> str:
        return "seccion_2_mesa_servicio.docx"
    
    def __init__(self, ):        
        self.glpi_extractor = None  # Se inicializa cuando se necesite (async)
        self.datos = {}
        self.doc = None
        self.anio = None
        self.mes = None
    
    def _construir_ruta_mes(self, mes: int, anio: int) -> str:
        """
        Construye la ruta en formato '01OCT – 31OCT' basándose en el mes y año.
        
        Args:
            mes: Número del mes (1-12)
            anio: Año (ej: 2025)
            
        Returns:
            String en formato '01OCT – 31OCT'
        """
        # Abreviaciones de meses en mayúsculas
        meses_abrev = {
            1: "ENE", 2: "FEB", 3: "MAR", 4: "ABR",
            5: "MAY", 6: "JUN", 7: "JUL", 8: "AGO",
            9: "SEP", 10: "OCT", 11: "NOV", 12: "DIC"
        }
        
        # Obtener el último día del mes
        ultimo_dia = calendar.monthrange(anio, mes)[1]
        
        # Construir la ruta
        mes_abrev = meses_abrev.get(mes, "")
        ruta = f"01{mes_abrev} – {ultimo_dia:02d}{mes_abrev}"
        
        return ruta
    
    def _construir_fecha_mes(self, mes: int, anio: int) -> str:
        """
        Construye la fecha en formato '01 al 31 de OCTUBRE DE 2025' basándose en el mes y año.
        
        Args:
            mes: Número del mes (1-12)
            anio: Año (ej: 2025)
            
        Returns:
            String en formato '01 al 31 de OCTUBRE DE 2025'
        """
        # Nombres de meses completos en mayúsculas
        meses_completos = {
            1: "ENERO", 2: "FEBRERO", 3: "MARZO", 4: "ABRIL",
            5: "MAYO", 6: "JUNIO", 7: "JULIO", 8: "AGOSTO",
            9: "SEPTIEMBRE", 10: "OCTUBRE", 11: "NOVIEMBRE", 12: "DICIEMBRE"
        }
        
        # Obtener el último día del mes
        ultimo_dia = calendar.monthrange(anio, mes)[1]
        
        # Obtener el nombre del mes en mayúsculas
        mes_nombre = meses_completos.get(mes, mes)
        
        # Construir la fecha en el formato solicitado
        fecha = f"01 al {ultimo_dia:02d} de {mes_nombre} DE {anio}"
        
        return fecha
    
    def _obtener_ultimo_dia_mes(self, mes: int, anio: int) -> str:
        """
        Obtiene el último día del mes en formato '30 de SEPTIEMBRE de 2025'.
        
        Args:
            mes: Número del mes (1-12)
            anio: Año (ej: 2025)
            
        Returns:
            String en formato '30 de SEPTIEMBRE de 2025'
        """
        # Nombres de meses completos en mayúsculas
        meses_completos = {
            1: "ENERO", 2: "FEBRERO", 3: "MARZO", 4: "ABRIL",
            5: "MAYO", 6: "JUNIO", 7: "JULIO", 8: "AGOSTO",
            9: "SEPTIEMBRE", 10: "OCTUBRE", 11: "NOVIEMBRE", 12: "DICIEMBRE"
        }
        
        # Obtener el último día del mes
        ultimo_dia = calendar.monthrange(anio, mes)[1]
        
        # Obtener el nombre del mes en mayúsculas
        mes_nombre = meses_completos.get(mes, "")
        
        # Construir la fecha en el formato solicitado
        fecha = f"{ultimo_dia} de {mes_nombre} de {anio}"
        
        return fecha

    def _seccion_2(self, data: Dict[str, Any]):       
     
        content_data = {
            "anio": data.get("anio"),
            "mes": data.get("mes"),
            "user_id": data.get("user_id"),
            "name_file": data.get("name_file"),
            "section_id": data.get("section_id"),
            "level": data.get("level"),
            "content": {
                "image": data.get("content").get("image"),                
            },
        }
        return content_data

    async def _seccion_2_1_mesa_servicio(self, data: Dict[str, Any]):
        
        from src.services.glpi_service import get_glpi_service
        glpi_service = await get_glpi_service()

        content = data.get("content", {})
        anio = data.get("anio")
        mes = data.get("mes")
        table_2 = await glpi_service.get_actividades_por_subsistema(anio, mes)            
        
        content_data = {
            "anio": anio,
            "mes": mes,
            "user_id": data.get("user_id"),
            "name_file": data.get("name_file"),
            "section_id": data.get("section_id"),
            "level": data.get("level"),
            "content": {
                "route": content.get("route", "") or (self._construir_ruta_mes(mes, anio) if mes and anio else ""),
                "image": content.get("image", ""),
                "table_1": content.get("table_1", []),
                "table_2": table_2,
            },
        }
        
        return content_data
    
    def _seccion_2_2_herramientas(self, data: Dict[str, Any]):
        content = data.get("content", {})
        content_data = {
            "anio": data.get("anio"),
            "mes": data.get("mes"),
            "user_id": data.get("user_id"),
            "name_file": data.get("name_file"),
            "section_id": data.get("section_id"),
            "level": data.get("level"),
            "content": {
                "email": content.get("email", ""),
            },
        }
        
        return content_data
    
    async def _seccion_2_3_visitas_diagnostico(self, data: Dict[str, Any]):
        content = data.get("content", {})
        anio = data.get("anio")
        mes = data.get("mes")
        
        # Obtener datos desde GLPI si no están en el content
        table_1 = content.get("table_1", [])
        if not table_1 and anio and mes:
            try:
                from src.services.glpi_service import get_glpi_service
                glpi_service = await get_glpi_service()
                table_1 = await glpi_service.get_visitas_diagnostico_subsistemas(anio, mes)
            except Exception as e:
                logger.warning(f"Error al obtener table_1 desde GLPI MySQL para sección 2.3: {e}. Usando datos de content.")
                table_1 = content.get("table_1", [])
        
        content_data = {
            "anio": anio,
            "mes": mes,
            "user_id": data.get("user_id"),
            "name_file": data.get("name_file"),
            "section_id": data.get("section_id"),
            "level": data.get("level"),
            "content": {
                "table_1": table_1,
                "comunicacion": content.get("comunicacion", ""),
                "oficio": content.get("oficio", ""),
            },
        }
        
        return content_data
    
    async def _seccion_2_4_tickets(self, data: Dict[str, Any]):
        content = data.get("content", {})
        anio = data.get("anio")
        mes = data.get("mes")

        from src.services.glpi_service import get_glpi_service
        glpi_service = await get_glpi_service()
        table_1 = await glpi_service.get_actividades_por_subsistema(anio, mes) 
        table_2 = await glpi_service.get_estado_tickets_por_subsistema(anio, mes)
        
        content_data = {
            "anio": anio,
            "mes": mes,
            "user_id": data.get("user_id"),
            "name_file": data.get("name_file"),
            "section_id": data.get("section_id"),
            "level": data.get("level"),
            "content": {
                "table_1": table_1,
                "name_document": content.get("name_document", ""),
                "table_2": table_2,
            },
        }
        
        return content_data
    
    async def _seccion_2_5_escalamientos(self, data: Dict[str, Any]):
        content = data.get("content", {})
        anio = data.get("anio")
        mes = data.get("mes")
        
        # Ruta del Excel definida dentro del método (no se solicita en el JSON)
        ruta_excel = "https://verytelcsp.sharepoint.com/sites/OPERACIONES/Shared%20Documents/PROYECTOS/A%C3%B1o%202024/2024-1809%20MANTTO%20BOGOTA%20ETB/8.%20INFORMES/INFORME%20MENSUAL/13.%2001NOV%20-%2030NOV/01%20OBLIGACIONES%20GENERALES/OBLIGACI%C3%93N%202,5,6,9,13/ANEXO%20MESA%20DE%20SERVICIO/ESCALAMIENTOS/ESCALAMIENTOS.xlsx"
        
        # Obtener datos desde SharePoint para calcular los totales
        table_1 = content.get("table_1", [])  # Fallback si falla SharePoint
        
        if ruta_excel:
            try:
                from src.services.sharepoint_service import get_sharepoint_service
                import config
                sharepoint_service = await get_sharepoint_service()
                
                # Convertir URL completa a ruta relativa del servidor si es necesario
                if ruta_excel.startswith("http://") or ruta_excel.startswith("https://"):
                    ruta_sharepoint = convertir_url_sharepoint_a_ruta_relativa(ruta_excel)
                    ruta_base = None  # No usar ruta_base cuando ya es ruta relativa del servidor
                else:
                    # Es una ruta relativa simple, usar ruta_base para construir la ruta completa
                    ruta_sharepoint = ruta_excel
                    ruta_base = getattr(config, 'SHAREPOINT_BASE_PATH', None) or os.getenv("SHAREPOINT_BASE_PATH")
                
                # Leer Excel completo con las 3 hojas
                resultado = await sharepoint_service.leer_excel_escalamientos_completo(
                    ruta_textual=ruta_sharepoint,
                    ruta_base=ruta_base
                )
                
                if resultado.get("success"):
                    # Obtener los datos de cada hoja
                    datos_enel = resultado.get("datos", {}).get("enel", [])
                    datos_caida_masiva = resultado.get("datos", {}).get("caida_masiva", [])
                    datos_conectividad = resultado.get("datos", {}).get("conectividad", [])
                    
                    # Calcular las cantidades (excluyendo filas vacías o con valores vacíos)
                    cantidad_enel = len([r for r in datos_enel if r and any(str(v).strip() for k, v in r.items() if k != 'tipo')])
                    cantidad_caida_masiva = len([r for r in datos_caida_masiva if r and any(str(v).strip() for k, v in r.items() if k != 'tipo')])
                    cantidad_conectividad = len([r for r in datos_conectividad if r and any(str(v).strip() for k, v in r.items() if k != 'tipo')])
                    
                    # Calcular el total
                    total = cantidad_enel + cantidad_caida_masiva + cantidad_conectividad
                    
                    # Construir la tabla de resumen
                    table_1 = [
                        {"escalamiento": "ENEL", "cantidad": str(cantidad_enel)},
                        {"escalamiento": "CAÍDA MASIVA", "cantidad": str(cantidad_caida_masiva)},
                        {"escalamiento": "CONECTIVIDAD", "cantidad": str(cantidad_conectividad)},
                        {"escalamiento": "TOTAL", "cantidad": str(total)}
                    ]
                    
                    logger.info(f"Datos de escalamientos calculados desde SharePoint: ENEL={cantidad_enel}, CAÍDA MASIVA={cantidad_caida_masiva}, CONECTIVIDAD={cantidad_conectividad}, TOTAL={total}")
                else:
                    logger.warning(f"Error al obtener datos desde SharePoint: {resultado.get('message')}. Usando datos de content como fallback.")
                    if not table_1:
                        # Si no hay datos de fallback, crear estructura vacía
                        table_1 = [
                            {"escalamiento": "ENEL", "cantidad": "0"},
                            {"escalamiento": "CAÍDA MASIVA", "cantidad": "0"},
                            {"escalamiento": "CONECTIVIDAD", "cantidad": "0"},
                            {"escalamiento": "TOTAL", "cantidad": "0"}
                        ]
            except Exception as e:
                logger.warning(f"Error al obtener table_1 desde SharePoint para sección 2.5: {e}. Usando datos de content como fallback.")
                if not table_1:
                    # Si no hay datos de fallback, crear estructura vacía
                    table_1 = [
                        {"escalamiento": "ENEL", "cantidad": "0"},
                        {"escalamiento": "CAÍDA MASIVA", "cantidad": "0"},
                        {"escalamiento": "CONECTIVIDAD", "cantidad": "0"},
                        {"escalamiento": "TOTAL", "cantidad": "0"}
                    ]
        else:
            # Si no hay ruta_excel, usar datos de content o crear estructura vacía
            if not table_1:
                table_1 = [
                    {"escalamiento": "ENEL", "cantidad": "0"},
                    {"escalamiento": "CAÍDA MASIVA", "cantidad": "0"},
                    {"escalamiento": "CONECTIVIDAD", "cantidad": "0"},
                    {"escalamiento": "TOTAL", "cantidad": "0"}
                ]
        
        content_data = {
            "anio": anio,
            "mes": mes,
            "user_id": data.get("user_id"),
            "name_file": data.get("name_file"),
            "section_id": data.get("section_id"),
            "level": data.get("level"),
            "content": {
                "table_1": table_1,
                "ruta_excel": ruta_excel,
            },
        }
        
        return content_data
    
    async def _seccion_2_5_1_enel(self, data: Dict[str, Any]):
        content = data.get("content", {})
        anio = data.get("anio")
        mes = data.get("mes")
        
        # Ruta del Excel definida dentro del método (no se solicita en el JSON)
        ruta_excel = "https://verytelcsp.sharepoint.com/sites/OPERACIONES/Shared%20Documents/PROYECTOS/A%C3%B1o%202024/2024-1809%20MANTTO%20BOGOTA%20ETB/8.%20INFORMES/INFORME%20MENSUAL/13.%2001NOV%20-%2030NOV/01%20OBLIGACIONES%20GENERALES/OBLIGACI%C3%93N%202,5,6,9,13/ANEXO%20MESA%20DE%20SERVICIO/ESCALAMIENTOS/ESCALAMIENTOS.xlsx"
        
        # Siempre obtener datos desde SharePoint (hoja 3 - ENEL)
        # Siguiendo el patrón de seccion_1_info_general.py
        table_1 = content.get("table_1", [])  # Fallback si falla SharePoint
        
        if ruta_excel:
            try:
                from src.services.sharepoint_service import get_sharepoint_service
                sharepoint_service = await get_sharepoint_service()
                
                # Convertir URL completa a ruta relativa del servidor si es necesario
                # Siguiendo el patrón de seccion_1_info_general.py
                if ruta_excel.startswith("http://") or ruta_excel.startswith("https://"):
                    ruta_sharepoint = convertir_url_sharepoint_a_ruta_relativa(ruta_excel)
                    ruta_base = None  # No usar ruta_base cuando ya es ruta relativa del servidor
                else:
                    # Es una ruta relativa simple, usar ruta_base para construir la ruta completa
                    ruta_sharepoint = ruta_excel
                    ruta_base = getattr(config, 'SHAREPOINT_BASE_PATH', None) or os.getenv("SHAREPOINT_BASE_PATH")
                
                # Leer Excel completo con las 3 hojas
                resultado = await sharepoint_service.leer_excel_escalamientos_completo(
                    ruta_textual=ruta_sharepoint,
                    ruta_base=ruta_base
                )
                
                if resultado.get("success"):
                    # Usar los datos de ENEL (hoja 3)
                    table_1 = resultado.get("datos", {}).get("enel", [])
                    # Formatear fechas en la tabla (campo: fecha_escalamiento)
                    table_1 = formatear_fechas_en_tabla(table_1, ["fecha_escalamiento", "fecha"])
                    logger.info(f"Datos de ENEL obtenidos desde SharePoint: {len(table_1)} registros")
                else:
                    logger.warning(f"Error al obtener datos desde SharePoint: {resultado.get('message')}. Usando datos de content como fallback.")
                    if not table_1:
                        table_1 = []
                    else:
                        # Formatear fechas también en el fallback
                        table_1 = formatear_fechas_en_tabla(table_1, ["fecha_escalamiento", "fecha"])
            except Exception as e:
                logger.warning(f"Error al obtener table_1 desde SharePoint para sección 2.5.1: {e}. Usando datos de content como fallback.")
                if not table_1:
                    table_1 = []
                else:
                    # Formatear fechas también en el fallback
                    table_1 = formatear_fechas_en_tabla(table_1, ["fecha_escalamiento", "fecha"])
        
        content_data = {
            "anio": anio,
            "mes": mes,
            "user_id": data.get("user_id"),
            "name_file": data.get("name_file"),
            "section_id": data.get("section_id"),
            "level": data.get("level"),
            "content": {
                "table_1": table_1,
                "ruta_excel": ruta_excel,
            },
        }
        
        return content_data
    
    async def _seccion_2_5_2_caida_masiva(self, data: Dict[str, Any]):
        content = data.get("content", {})
        anio = data.get("anio")
        mes = data.get("mes")
        
        # Ruta del Excel definida dentro del método (no se solicita en el JSON)
        ruta_excel = "https://verytelcsp.sharepoint.com/sites/OPERACIONES/Shared%20Documents/PROYECTOS/A%C3%B1o%202024/2024-1809%20MANTTO%20BOGOTA%20ETB/8.%20INFORMES/INFORME%20MENSUAL/13.%2001NOV%20-%2030NOV/01%20OBLIGACIONES%20GENERALES/OBLIGACI%C3%93N%202,5,6,9,13/ANEXO%20MESA%20DE%20SERVICIO/ESCALAMIENTOS/ESCALAMIENTOS.xlsx"
        
        # Siempre obtener datos desde SharePoint (hoja 1 - CAÍDA MASIVA)
        # Siguiendo el patrón de seccion_1_info_general.py
        table_1 = content.get("table_1", [])  # Fallback si falla SharePoint
        
        if ruta_excel:
            try:
                from src.services.sharepoint_service import get_sharepoint_service
                sharepoint_service = await get_sharepoint_service()
                
                # Convertir URL completa a ruta relativa del servidor si es necesario
                # Siguiendo el patrón de seccion_1_info_general.py
                if ruta_excel.startswith("http://") or ruta_excel.startswith("https://"):
                    ruta_sharepoint = convertir_url_sharepoint_a_ruta_relativa(ruta_excel)
                    ruta_base = None  # No usar ruta_base cuando ya es ruta relativa del servidor
                else:
                    # Si es una ruta relativa simple, usar ruta_base para construir la ruta completa
                    ruta_sharepoint = ruta_excel
                    ruta_base = getattr(config, 'SHAREPOINT_BASE_PATH', None) or os.getenv("SHAREPOINT_BASE_PATH")
                
                # Leer Excel completo con las 3 hojas
                resultado = await sharepoint_service.leer_excel_escalamientos_completo(
                    ruta_textual=ruta_sharepoint,
                    ruta_base=ruta_base
                )
                
                if resultado.get("success"):
                    # Usar los datos de CAÍDA MASIVA (hoja 1)
                    table_1 = resultado.get("datos", {}).get("caida_masiva", [])
                    # Formatear fechas en la tabla (campo: fecha)
                    table_1 = formatear_fechas_en_tabla(table_1, ["fecha", "fecha_escalamiento"])
                    logger.info(f"Datos de CAÍDA MASIVA obtenidos desde SharePoint: {len(table_1)} registros")
                else:
                    logger.warning(f"Error al obtener datos desde SharePoint: {resultado.get('message')}. Usando datos de content como fallback.")
                    if not table_1:
                        table_1 = []
                    else:
                        # Formatear fechas también en el fallback
                        table_1 = formatear_fechas_en_tabla(table_1, ["fecha", "fecha_escalamiento"])
            except Exception as e:
                logger.warning(f"Error al obtener table_1 desde SharePoint para sección 2.5.2: {e}. Usando datos de content como fallback.")
                if not table_1:
                    table_1 = []
                else:
                    # Formatear fechas también en el fallback
                    table_1 = formatear_fechas_en_tabla(table_1, ["fecha", "fecha_escalamiento"])
        
        content_data = {
            "anio": anio,
            "mes": mes,
            "user_id": data.get("user_id"),
            "name_file": data.get("name_file"),
            "section_id": data.get("section_id"),
            "level": data.get("level"),
            "content": {
                "table_1": table_1,
                "ruta_excel": ruta_excel,
            },
        }
        
        return content_data
    
    async def _seccion_2_5_3_conectividad(self, data: Dict[str, Any]):
        content = data.get("content", {})
        anio = data.get("anio")
        mes = data.get("mes")
        
        # Ruta del Excel definida dentro del método (no se solicita en el JSON)
        ruta_excel = "https://verytelcsp.sharepoint.com/sites/OPERACIONES/Shared%20Documents/PROYECTOS/A%C3%B1o%202024/2024-1809%20MANTTO%20BOGOTA%20ETB/8.%20INFORMES/INFORME%20MENSUAL/13.%2001NOV%20-%2030NOV/01%20OBLIGACIONES%20GENERALES/OBLIGACI%C3%93N%202,5,6,9,13/ANEXO%20MESA%20DE%20SERVICIO/ESCALAMIENTOS/ESCALAMIENTOS.xlsx"
        
        # Siempre obtener datos desde SharePoint (hoja 2 - CONECTIVIDAD)
        # Siguiendo el patrón de seccion_1_info_general.py
        table_1 = content.get("table_1", [])  # Fallback si falla SharePoint
        
        if ruta_excel:
            try:
                from src.services.sharepoint_service import get_sharepoint_service
                sharepoint_service = await get_sharepoint_service()
                
                # Convertir URL completa a ruta relativa del servidor si es necesario
                # Siguiendo el patrón de seccion_1_info_general.py
                if ruta_excel.startswith("http://") or ruta_excel.startswith("https://"):
                    ruta_sharepoint = convertir_url_sharepoint_a_ruta_relativa(ruta_excel)
                    ruta_base = None  # No usar ruta_base cuando ya es ruta relativa del servidor
                else:
                    # Si es una ruta relativa simple, usar ruta_base para construir la ruta completa
                    ruta_sharepoint = ruta_excel
                    ruta_base = getattr(config, 'SHAREPOINT_BASE_PATH', None) or os.getenv("SHAREPOINT_BASE_PATH")
                
                # Leer Excel completo con las 3 hojas
                resultado = await sharepoint_service.leer_excel_escalamientos_completo(
                    ruta_textual=ruta_sharepoint,
                    ruta_base=ruta_base
                )
                
                if resultado.get("success"):
                    # Usar los datos de CONECTIVIDAD (hoja 2)
                    table_1 = resultado.get("datos", {}).get("conectividad", [])
                    # Formatear fechas en la tabla (campo: fecha_escalamiento)
                    table_1 = formatear_fechas_en_tabla(table_1, ["fecha_escalamiento", "fecha"])
                    logger.info(f"Datos de CONECTIVIDAD obtenidos desde SharePoint: {len(table_1)} registros")
                else:
                    logger.warning(f"Error al obtener datos desde SharePoint: {resultado.get('message')}. Usando datos de content como fallback.")
                    if not table_1:
                        table_1 = []
                    else:
                        # Formatear fechas también en el fallback
                        table_1 = formatear_fechas_en_tabla(table_1, ["fecha_escalamiento", "fecha"])
            except Exception as e:
                logger.warning(f"Error al obtener table_1 desde SharePoint para sección 2.5.3: {e}. Usando datos de content como fallback.")
                if not table_1:
                    table_1 = []
                else:
                    # Formatear fechas también en el fallback
                    table_1 = formatear_fechas_en_tabla(table_1, ["fecha_escalamiento", "fecha"])
        
        content_data = {
            "anio": anio,
            "mes": mes,
            "user_id": data.get("user_id"),
            "name_file": data.get("name_file"),
            "section_id": data.get("section_id"),
            "level": data.get("level"),
            "content": {
                "table_1": table_1,
                "ruta_excel": ruta_excel,
            },
        }
        
        return content_data
    
    async def _seccion_2_6_hojas_vida(self, data: Dict[str, Any]):
        content = data.get("content", {})
        anio = data.get("anio")
        mes = data.get("mes")
        
        # Ruta de la carpeta en SharePoint donde buscar el Excel
        ruta_carpeta_sharepoint = "https://verytelcsp.sharepoint.com/sites/OPERACIONES/Shared%20Documents/PROYECTOS/A%C3%B1o%202024/2024-1809%20MANTTO%20BOGOTA%20ETB/8.%20INFORMES/INFORME%20MENSUAL/13.%2001NOV%20-%2030NOV/01%20OBLIGACIONES%20GENERALES/OBLIGACI%C3%93N%202,5,6,9,13/ANEXO%20MESA%20DE%20SERVICIO/ESTADO%20DEL%20SISTEMA"
        
        # Valor por defecto
        name_document = content.get("name_document", "NO EXISTE")
        
        # Buscar archivo Excel en la carpeta de SharePoint
        try:
            from src.services.sharepoint_service import get_sharepoint_service
            sharepoint_service = await get_sharepoint_service()
            
            # Convertir URL completa a ruta relativa al drive root
            # Necesitamos extraer solo la parte que viene después del base_path
            if ruta_carpeta_sharepoint.startswith("http://") or ruta_carpeta_sharepoint.startswith("https://"):
                from urllib.parse import unquote, urlparse
                import config
                url_parsed = urlparse(ruta_carpeta_sharepoint)
                path_decodificado = unquote(url_parsed.path)
                path_parts = [p for p in path_decodificado.split('/') if p]
                
                # Buscar "Shared Documents" en la ruta
                try:
                    idx_shared_docs = next(i for i, part in enumerate(path_parts) if part == "Shared Documents")
                    # Tomar todo después de "Shared Documents"
                    ruta_completa_desde_shared = '/'.join(path_parts[idx_shared_docs + 1:])
                    
                    # Obtener el base_path configurado (sin "Shared Documents")
                    base_path_config = getattr(config, 'SHAREPOINT_BASE_PATH', None) or os.getenv("SHAREPOINT_BASE_PATH", "")
                    if base_path_config:
                        # Normalizar base_path: remover "Shared Documents" si está al inicio
                        base_normalizado = base_path_config.rstrip('/').rstrip(' ')
                        if base_normalizado.startswith('Shared Documents/'):
                            base_normalizado = base_normalizado[len('Shared Documents/'):]
                        elif base_normalizado.startswith('Shared Documents'):
                            base_normalizado = base_normalizado[len('Shared Documents'):].lstrip('/')
                        
                        # Si la ruta completa empieza con el base_path, removerlo
                        if ruta_completa_desde_shared.startswith(base_normalizado):
                            ruta_carpeta = ruta_completa_desde_shared[len(base_normalizado):].lstrip('/')
                            logger.info(f"Base path removido. Ruta resultante: {ruta_carpeta}")
                        else:
                            # Si no coincide, usar la ruta completa después de Shared Documents
                            ruta_carpeta = ruta_completa_desde_shared
                            logger.info(f"Base path no coincide. Usando ruta completa: {ruta_carpeta}")
                    else:
                        # Si no hay base_path, usar toda la ruta después de Shared Documents
                        ruta_carpeta = ruta_completa_desde_shared
                        logger.info(f"No hay base_path configurado. Usando ruta completa: {ruta_carpeta}")
                        
                except StopIteration:
                    # Si no encuentra "Shared Documents", usar la ruta completa después de "sites"
                    try:
                        idx_sites = next(i for i, part in enumerate(path_parts) if part in ['sites', 'teams', 'personal'])
                        ruta_carpeta = '/'.join(path_parts[idx_sites + 2:])  # Saltar sites y nombre del sitio
                    except StopIteration:
                        ruta_carpeta = '/'.join(path_parts)
            else:
                ruta_carpeta = ruta_carpeta_sharepoint
            
            logger.info(f"Ruta carpeta para listar archivos: {ruta_carpeta}")
            
            # Listar archivos en la carpeta
            archivos = sharepoint_service.extractor.listar_archivos_en_carpeta(ruta_carpeta)
            
            logger.info(f"Archivos encontrados en la carpeta: {len(archivos) if archivos else 0}")
            if archivos:
                logger.info(f"Nombres de archivos encontrados: {[a.get('nombre', 'N/A') for a in archivos[:5]]}")
            
            if archivos:
                # Buscar archivos Excel (.xlsx, .xls)
                archivos_excel = [
                    archivo for archivo in archivos
                    if archivo.get("nombre", "").lower().endswith((".xlsx", ".xls"))
                ]
                
                if archivos_excel:
                    # Tomar el primer archivo Excel encontrado
                    nombre_archivo = archivos_excel[0].get("nombre", "")
                    if nombre_archivo:
                        name_document = nombre_archivo
                        logger.info(f"Archivo Excel encontrado en SharePoint: {name_document}")
                    else:
                        name_document = "NO EXISTE"
                        logger.warning("No se encontró nombre de archivo Excel en la respuesta de SharePoint")
                else:
                    name_document = "NO EXISTE"
                    logger.warning(f"No se encontraron archivos Excel en la carpeta: {ruta_carpeta}")
            else:
                name_document = "NO EXISTE"
                logger.warning(f"No se encontraron archivos en la carpeta: {ruta_carpeta}")
                
        except Exception as e:
            logger.warning(f"Error al buscar archivo Excel en SharePoint para sección 2.6: {e}. Usando valor por defecto.")
            name_document = content.get("name_document", "NO EXISTE")
        
        content_data = {
            "anio": anio,
            "mes": mes,
            "user_id": data.get("user_id"),
            "name_file": data.get("name_file"),
            "section_id": data.get("section_id"),
            "level": data.get("level"),
            "content": {
                "name_document": name_document,
            },
        }
        
        return content_data
    
    def _seccion_2_7_estado_sistema(self, data: Dict[str, Any]):
        content = data.get("content", {})
        anio = data.get("anio")
        mes = data.get("mes")
        
        # Estados estáticos para la tabla 1
        estados_estaticos = [
            "CAÍDA MASIVA",
            "FUERA DE SERVICIO",
            "OPERATIVA",
            "OPERATIVA CON NOVEDAD",
            "TOTAL"
        ]
        
        # Si table_1 ya viene en el content, usarla directamente
        # Si no, construirla desde cantidades_estado
        table_1 = content.get("table_1", [])
        
        # Si table_1 está vacía, construirla con estados estáticos y cantidades dinámicas
        if not table_1:
            # Obtener cantidades del content
            cantidades = content.get("cantidades_estado", {})
            
            # Construir tabla_1 con estados estáticos y cantidades dinámicas
            total_calculado = 0
            
            for estado in estados_estaticos:
                if estado == "TOTAL":
                    # Calcular el total sumando las cantidades anteriores
                    cantidad_total = cantidades.get("total", str(total_calculado))
                    # Si no se proporciona total, usar el calculado
                    if not cantidad_total or cantidad_total == "":
                        cantidad_total = str(total_calculado)
                    table_1.append({
                        "estado": estado,
                        "cantidad": str(cantidad_total)
                    })
                else:
                    # Obtener la cantidad para este estado
                    # Mapear nombres de estados a claves posibles en el JSON
                    clave_estado = estado.lower().replace(" ", "_").replace("í", "i")
                    cantidad = cantidades.get(clave_estado, "")
                    
                    # Si no se encuentra con la clave normalizada, buscar variaciones
                    if not cantidad or cantidad == "":
                        # Buscar en diferentes formatos
                        variaciones = [
                            estado.lower().replace(" ", "_"),
                            estado.lower().replace(" ", ""),
                            estado.upper().replace(" ", "_"),
                            estado.upper().replace(" ", ""),
                        ]
                        for var in variaciones:
                            if var in cantidades:
                                cantidad = cantidades[var]
                                break
                    
                    # Si aún no hay cantidad, usar vacío
                    if not cantidad:
                        cantidad = ""
                    else:
                        # Sumar al total si es un número válido
                        try:
                            total_calculado += int(str(cantidad).replace(",", "").strip())
                        except (ValueError, TypeError):
                            pass
                    
                    table_1.append({
                        "estado": estado,
                        "cantidad": str(cantidad)
                    })
        
        # Responsables estáticos para la tabla 2
        responsables_estaticos = [
            "PTE APROBACIÓN USO DE BOLSA",
            "CONECTIVIDAD",
            "MANTENIMIENTO",
            "SINIESTRO",
            "ENERGIZACIÓN",
            "OBRAS",
            "PUNTO DESMONTADO",
            "ENEL",
            "TOTAL"
        ]
        
        # Si table_2 ya viene en el content, usarla directamente
        # Si no, construirla desde cantidades_responsable
        table_2 = content.get("table_2", [])
        
        # Si table_2 está vacía, construirla con responsables estáticos y cantidades dinámicas
        if not table_2:
            # Obtener cantidades del content
            cantidades_responsable = content.get("cantidades_responsable", {})
            
            # Construir table_2 con responsables estáticos y cantidades dinámicas
            total_calculado_responsable = 0
            
            for responsable in responsables_estaticos:
                if responsable == "TOTAL":
                    # Calcular el total sumando las cantidades anteriores
                    cantidad_total_responsable = cantidades_responsable.get("total", str(total_calculado_responsable))
                    # Si no se proporciona total, usar el calculado
                    if not cantidad_total_responsable or cantidad_total_responsable == "":
                        cantidad_total_responsable = str(total_calculado_responsable)
                    table_2.append({
                        "responsable": responsable,
                        "cantidad": str(cantidad_total_responsable)
                    })
                else:
                    # Obtener la cantidad para este responsable
                    # Mapear nombres de responsables a claves posibles en el JSON
                    clave_responsable = responsable.lower().replace(" ", "_").replace("ó", "o").replace("á", "a")
                    cantidad = cantidades_responsable.get(clave_responsable, "")
                    
                    # Si no se encuentra con la clave normalizada, buscar variaciones
                    if not cantidad or cantidad == "":
                        # Buscar en diferentes formatos
                        variaciones = [
                            responsable.lower().replace(" ", "_"),
                            responsable.lower().replace(" ", ""),
                            responsable.upper().replace(" ", "_"),
                            responsable.upper().replace(" ", ""),
                            "pte_aprobacion_uso_de_bolsa" if "PTE APROBACIÓN" in responsable else None,
                            "pte_aprobacion_uso_de_bolsa".upper() if "PTE APROBACIÓN" in responsable else None,
                        ]
                        for var in variaciones:
                            if var and var in cantidades_responsable:
                                cantidad = cantidades_responsable[var]
                                break
                    
                    # Si aún no hay cantidad, usar vacío
                    if not cantidad:
                        cantidad = ""
                    else:
                        # Sumar al total si es un número válido
                        try:
                            total_calculado_responsable += int(str(cantidad).replace(",", "").strip())
                        except (ValueError, TypeError):
                            pass
                    
                    table_2.append({
                        "responsable": responsable,
                        "cantidad": str(cantidad)
                    })
        
        # Subsistemas estáticos para la tabla 3
        subsistemas_estaticos = [
            "ESTACIONES DE POLICÍA",
            "PROYECTO 350",
            "PROYECTO 732",
            "PROYECTO ALCALDÍA",
            "PROYECTO CAI",
            "PROYECTO COLEGIOS",
            "PROYECTO CTP",
            "PROYECTO ESU-C4",
            "PROYECTO ESU-ESTADIO",
            "PROYECTO FVS",
            "PROYECTO TRANSMILENIO",
            "TOTAL"
        ]
        
        # Columnas dinámicas para la tabla 3
        columnas_dinamicas = [
            "caida_masiva",
            "fuera_de_servicio",
            "operativa",
            "operativa_con_novedad",
            "total"
        ]
        
        # Si table_3 ya viene en el content, usarla directamente
        # Si no, construirla desde datos_subsistemas
        table_3 = content.get("table_3", [])
        
        # Si table_3 está vacía, construirla con subsistemas estáticos y datos dinámicos
        if not table_3:
            # Obtener datos de subsistemas del content
            datos_subsistemas = content.get("datos_subsistemas", {})
            
            # Construir table_3 con subsistemas estáticos y datos dinámicos
            totales_columnas = {
                "caida_masiva": 0,
                "fuera_de_servicio": 0,
                "operativa": 0,
                "operativa_con_novedad": 0,
                "total": 0
            }
            
            for subsistema in subsistemas_estaticos:
                fila = {"subsistema": subsistema}
                
                if subsistema == "TOTAL":
                    # Para la fila TOTAL, usar los totales calculados
                    for columna in columnas_dinamicas:
                        valor_total = datos_subsistemas.get("total", {}).get(columna, str(totales_columnas[columna]))
                        if not valor_total or valor_total == "":
                            valor_total = str(totales_columnas[columna])
                        fila[columna] = str(valor_total)
                else:
                    # Obtener datos para este subsistema
                    # Normalizar nombre del subsistema para buscar en datos_subsistemas
                    clave_subsistema = subsistema.lower().replace(" ", "_").replace("ó", "o").replace("á", "a").replace("í", "i")
                    
                    datos_subsistema = datos_subsistemas.get(clave_subsistema, {})
                    
                    # Si no se encuentra con la clave normalizada, buscar variaciones
                    if not datos_subsistema:
                        variaciones = [
                            subsistema.lower().replace(" ", "_"),
                            subsistema.lower().replace(" ", ""),
                            subsistema.upper().replace(" ", "_"),
                            subsistema.upper().replace(" ", ""),
                        ]
                        for var in variaciones:
                            if var in datos_subsistemas:
                                datos_subsistema = datos_subsistemas[var]
                                break
                    
                    # Llenar cada columna dinámica
                    for columna in columnas_dinamicas:
                        valor = datos_subsistema.get(columna, "")
                        
                        # Si no se encuentra, buscar variaciones del nombre de la columna
                        if not valor or valor == "":
                            variaciones_columna = [
                                columna,
                                columna.replace("_", " "),
                                columna.upper(),
                                columna.upper().replace("_", " "),
                            ]
                            for var_col in variaciones_columna:
                                if var_col in datos_subsistema:
                                    valor = datos_subsistema[var_col]
                                    break
                        
                        # Si aún no hay valor, usar vacío
                        if not valor:
                            valor = ""
                    else:
                        # Sumar al total de la columna si es un número válido (excepto para la columna total)
                        if columna != "total":
                            try:
                                totales_columnas[columna] += int(str(valor).replace(",", "").strip())
                            except (ValueError, TypeError):
                                pass
                    
                    fila[columna] = str(valor)
                    
                    # Calcular el total de la fila si no se proporcionó
                    if not fila.get("total") or fila.get("total") == "":
                        try:
                            total_fila = (
                                int(str(fila.get("caida_masiva", "0")).replace(",", "").strip() or "0") +
                                int(str(fila.get("fuera_de_servicio", "0")).replace(",", "").strip() or "0") +
                                int(str(fila.get("operativa", "0")).replace(",", "").strip() or "0") +
                                int(str(fila.get("operativa_con_novedad", "0")).replace(",", "").strip() or "0")
                            )
                            fila["total"] = str(total_fila)
                            totales_columnas["total"] += total_fila
                        except (ValueError, TypeError):
                            fila["total"] = ""
                
                table_3.append(fila)
        
        content_data = {
            "anio": anio,
            "mes": mes,
            "user_id": data.get("user_id"),
            "name_file": data.get("name_file"),
            "section_id": data.get("section_id"),
            "level": data.get("level"),
            "content": {
                "table_1": table_1,
                "image": content.get("image", ""),
                "section_1": content.get("section_1", ""),
                "section_2": content.get("section_2", ""),
                "table_2": table_2,
                "section_3": content.get("section_3", ""),
                "table_3": table_3,
                "name_document": content.get("name_document", ""),
                "observaciones": content.get("observaciones", ""),
            },
        }
        
        return content_data
    
    
    def _groupby_consecutivo_caida(self, table_data: List[Dict[str, Any]], campo: str = "consecutivo_caida") -> Dict[str, int]:
        """
        Agrupa los registros por el campo 'consecutivo_caida' y cuenta cuántos registros hay para cada consecutivo.
        
        Args:
            table_data: Lista de diccionarios con los datos de la tabla
            campo: Nombre del campo por el cual agrupar (por defecto "consecutivo_caida")
            
        Returns:
            Diccionario con el consecutivo como clave y la cantidad de registros como valor
            Ejemplo: {"CM-001": 5, "CM-002": 3, "CM-003": 2}
        """
        if not table_data:
            return {}
        
        # Diccionario para almacenar el conteo por consecutivo
        conteo_por_consecutivo = {}
        
        # Buscar el campo en diferentes variaciones (case-insensitive)
        campos_posibles = [
            campo.lower(),
            campo.upper(),
            campo.title(),
            "consecutivo_caida",
            "CONSECUTIVO CAÍDA",
            "Consecutivo Caída",
            "consecutivo caida",
            "CONSECUTIVO_CAIDA",
            "consecutivo_caida",
        ]
        
        for registro in table_data:
            if not registro:
                continue
            
            # Buscar el valor del consecutivo en el registro
            consecutivo = None
            for campo_posible in campos_posibles:
                if campo_posible in registro:
                    valor = registro[campo_posible]
                    if valor and str(valor).strip():
                        consecutivo = str(valor).strip()
                break
        
            # Si no se encuentra con los nombres estándar, buscar cualquier campo que contenga "consecutivo"
            if not consecutivo:
                for key, value in registro.items():
                    if "consecutivo" in key.lower() and value and str(value).strip():
                        consecutivo = str(value).strip()
                        break
            
            # Si aún no se encuentra, usar "SIN CONSECUTIVO"
            if not consecutivo:
                consecutivo = "SIN CONSECUTIVO"
            
            # Contar el registro
            conteo_por_consecutivo[consecutivo] = conteo_por_consecutivo.get(consecutivo, 0) + 1
        
        return conteo_por_consecutivo
    
    def _reemplazar_placeholder_con_tabla(self, doc: Document, placeholder: str, table_data: List[List[str]]):
        """
        Busca un placeholder en el documento y lo reemplaza con una tabla creada programáticamente.
        Wrapper para mantener compatibilidad con código existente.
        
        Args:
            doc: Documento Word (después de renderizar)
            placeholder: Texto del placeholder a buscar (ej: "[[TABLE_21_1]]")
            table_data: Lista de listas con los datos de la tabla (primera fila son headers)
        """
        reemplazar_placeholder_con_tabla(doc, placeholder, table_data)
    
    def _construir_contexto(self, document: Dict[str, Any], template: DocxTemplate) -> Dict[str, Any]:
        mes = document.get("mes")
        anio = document.get("anio")
        
        # Asegurar que mes y anio sean válidos
        mes_nombre = config.MESES.get(mes, "") if mes else ""

        index = document.get("index", [])
        if not isinstance(index, list):
            index = []
        
        # Construir la ruta del mes si no existe
        route_21 = ""
        if mes and anio:
            route_21 = self._construir_ruta_mes(mes, anio)
        
        # Construir la fecha en formato '01 al 31 de OCTUBRE DE 2025'
        date_month = ""
        if mes and anio:
            date_month = self._construir_fecha_mes(mes, anio)
        
        # Inicializar contexto base
        contexto = {
            "mes": mes_nombre,
            "anio": anio,
            "mes_numero": mes,
            "date_month": date_month
        }
        
        # Sección 2 (index[0]) - INFORME DE MESA DE SERVICIO
        if len(index) > 0:
            content_2 = index[0].get("content", {})
            contexto["image_2"] = procesar_imagen(template, content_2.get("image", "")) or ""
        
        # Sección 2.1 (index[1]) - INFORME DE MESA DE SERVICIO
        if len(index) > 1:
            content_21 = index[1].get("content", {})
            # Si la ruta está vacía en el documento, usar la construida
            route_doc = content_21.get("route", "")
            contexto["route_21"] = route_doc if route_doc else route_21
            contexto["image_21"] = procesar_imagen(template, content_21.get("image", "")) or ""
            
            # Preparar datos de tablas en formato lista de listas
            table_1_data = content_21.get("table_1", [])
            if table_1_data:
                headers_21_1 = ["ÍTEM", "FECHA", "REFERENCIA", "RADICADO", "ESTADO", "APROBACIÓN"]
                campos_21_1 = ["item", "fecha", "referencia", "radicado", "estado", "aprobacion"]
                contexto["table_21_1"] = preparar_datos_tabla(table_1_data, headers_21_1, campos_21_1)
            else:
                contexto["table_21_1"] = []
            
            table_2_data = content_21.get("table_2", [])
            if table_2_data:
                headers_21_2 = ["SUBSISTEMA", "DIAGNÓSTICO", "DIAGNÓSTICO SUBSISTEMA", "LIMPIEZA ACRÍLICO", 
                                "MTO ACOMETIDA", "MTO CORRECTIVO", "MTO CORRECTIVO SUBSISTEMA", "PLAN DE CHOQUE", "TOTAL"]
                campos_21_2 = ["subsistema", "diagnostico", "diagnostico_subsistema", "limpieza_acrilico",
                              "mto_acometida", "mto_correctivo", "mto_correctivo_subsistema", "plan_de_choque", "total"]
                
                # Calcular el total de diagnósticos (suma de la columna "diagnostico")
                total_diagnostico = 0
                try:
                    for item in table_2_data:
                        valor = item.get("diagnostico", "")
                        if valor:
                            if isinstance(valor, str):
                                valor_limpio = valor.replace(",", "").replace(" ", "").strip()
                                if valor_limpio:
                                    total_diagnostico += float(valor_limpio)
                            else:
                                total_diagnostico += float(valor)
                except (ValueError, TypeError):
                    pass
                
                # Asignar el total de diagnósticos al contexto
                contexto["diagnostico_21"] = str(int(total_diagnostico)) if total_diagnostico == int(total_diagnostico) else str(total_diagnostico)
                
                # Calcular el total general (suma de la columna "total")
                total_general = 0
                try:
                    for item in table_2_data:
                        valor = item.get("total", "")
                        if valor:
                            if isinstance(valor, str):
                                valor_limpio = valor.replace(",", "").replace(" ", "").strip()
                                if valor_limpio:
                                    total_general += float(valor_limpio)
                            else:
                                total_general += float(valor)
                except (ValueError, TypeError):
                    pass
                
                # Guardar el total general para reutilizarlo en total_24
                contexto["_total_general_table_2"] = int(total_general) if total_general == int(total_general) else total_general
                
                # Encontrar el subsistema con el mayor total
                mayor_subsistema = None
                mayor_total = 0
                try:
                    for item in table_2_data:
                        valor_total = item.get("total", "")
                        if valor_total:
                            # Convertir a número
                            if isinstance(valor_total, str):
                                valor_limpio = valor_total.replace(",", "").replace(" ", "").strip()
                                if valor_limpio:
                                    total_num = float(valor_limpio)
                                else:
                                    continue
                            else:
                                total_num = float(valor_total)
                            
                            # Comparar con el mayor encontrado hasta ahora
                            if total_num > mayor_total:
                                mayor_total = total_num
                                nombre_subsistema = item.get("subsistema", "")
                                mayor_subsistema = {
                                    "nombre": nombre_subsistema if nombre_subsistema else "N/A",
                                    "total": int(total_num) if total_num == int(total_num) else total_num
                                }
                except (ValueError, TypeError) as e:
                    logger.warning(f"Error al encontrar mayor subsistema: {e}")
                
                # Guardar el mayor subsistema para reutilizarlo
                contexto["_mayor_subsistema_table_2"] = mayor_subsistema
                
                # Agregar fila de totales (columna 0 es SUBSISTEMA donde pondremos "TOTAL")
                table_21_2_preparada = preparar_datos_tabla(
                    table_2_data, headers_21_2, campos_21_2, 
                    agregar_totales=True, columna_total_texto=0
                )
                contexto["table_21_2"] = table_21_2_preparada
                # Guardar table_2_data preparada para reutilizarla en table_24_1
                contexto["_table_2_data_preparada"] = table_21_2_preparada
            else:
                contexto["table_21_2"] = []
                contexto["diagnostico_21"] = "0"
                contexto["_table_2_data_preparada"] = []
                contexto["_total_general_table_2"] = 0
                contexto["_mayor_subsistema_table_2"] = None            
        # Sección 2.2 (index[2]) - HERRAMIENTAS DE TRABAJO
        if len(index) > 2:
            content_22 = index[2].get("content", {})
            contexto["email_22"] = content_22.get("email", "")
        
        # Sección 2.3 (index[3]) - VISITAS DE DIAGNÓSTICOS A SUBSISTEMAS
        if len(index) > 3:
            content_23 = index[3].get("content", {})
            contexto["comunicacion_23"] = content_23.get("comunicacion", "")
            contexto["oficio_23"] = content_23.get("oficio", "")
            
            # Preparar datos de tabla en formato lista de listas
            table_23_data = content_23.get("table_1", [])
            if table_23_data:
                headers_23_1 = ["SUBSISTEMA", "EJECUTADAS"]
                campos_23_1 = ["subsistema", "ejecutadas"]
                # Agregar fila de totales (columna 0 es SUBSISTEMA donde pondremos "TOTAL")
                contexto["table_23_1"] = preparar_datos_tabla(
                    table_23_data, headers_23_1, campos_23_1,
                    agregar_totales=True, columna_total_texto=0
                )
            else:
                contexto["table_23_1"] = []
        
        # Sección 2.4 (index[4]) - INFORME CONSOLIDADO DEL ESTADO DE LOS TICKETS ADMINISTRATIVOS
        if len(index) > 4:
            content_24 = index[4].get("content", {})
            # table_24_1 usa los mismos datos que table_21_2 (table_2_data preparada)
            contexto["table_24_1"] = contexto.get("_table_2_data_preparada", [])
            # total_24 es el total general de la tabla (suma de la columna "total")
            contexto["total_24"] = contexto.get("_total_general_table_2", 0)
            # mayor_subsistema_21: subsistema con mayor total en formato "nombre total con total tickets"
            mayor_subsistema_info = contexto.get("_mayor_subsistema_table_2")
            if mayor_subsistema_info:
                nombre = mayor_subsistema_info["nombre"]
                total = mayor_subsistema_info["total"]
                contexto["mayor_subsistema_21"] = f"{nombre} {total} con {total} tickets"
            else:
                contexto["mayor_subsistema_21"] = ""
            contexto["name_document_24"] = content_24.get("name_document", "")
            
            # Preparar datos de tabla table_24_2 en formato lista de listas
            table_24_2_data = content_24.get("table_2", [])
            if table_24_2_data:
                headers_24_2 = ["SUBSISTEMAS", "CERRADO", "EN CURSO (ASIGNADA)", "EN CURSO (PLANIFICADA)", 
                               "EN ESPERA", "RESUELTAS", "TOTAL"]
                campos_24_2 = ["subsistema", "cerrado", "en_curso_asignada", "en_curso_planificada",
                              "en_espera", "resueltas", "total"]
                # Agregar fila de totales (columna 0 es SUBSISTEMAS donde pondremos "TOTAL")
                contexto["table_24_2"] = preparar_datos_tabla(
                    table_24_2_data, headers_24_2, campos_24_2,
                    agregar_totales=True, columna_total_texto=0
                )
            else:
                contexto["table_24_2"] = []
        
        # Sección 2.5 (index[5]) - ESCALAMIENTOS
        if len(index) > 5:
            content_25 = index[5].get("content", {})
            table_25_data = content_25.get("table_1", [])
            if table_25_data:
                # Detectar columnas automáticamente desde los datos
                headers_25, campos_25 = detectar_columnas_desde_datos(table_25_data)
                contexto["table_25_1"] = preparar_datos_tabla(
                    table_25_data, headers_25, campos_25,
                    agregar_totales=False, columna_total_texto=0
                )
            else:
                contexto["table_25_1"] = []
        
        # Sección 2.5.1 (index[6]) - ENEL
        if len(index) > 6:
            content_251 = index[6].get("content", {})
            table_251_data = content_251.get("table_1", [])
            if table_251_data:
                # Detectar columnas automáticamente desde los datos
                headers_251, campos_251 = detectar_columnas_desde_datos(table_251_data)
                table_251_preparada = preparar_datos_tabla(
                    table_251_data, headers_251, campos_251,
                    agregar_totales=False, columna_total_texto=0
                )
                contexto["table_251_1"] = table_251_preparada
                # Contar registros de la tabla
                contexto["count_rows_251_1"] = count_rows(table_251_preparada)
            else:
                contexto["table_251_1"] = []
                contexto["count_rows_251_1"] = 0
        
        # Sección 2.5.2 (index[7]) - CAÍDA MASIVA
        if len(index) > 7:
            content_252 = index[7].get("content", {})
            table_252_data = content_252.get("table_1", [])
            if table_252_data:
                # Detectar columnas automáticamente desde los datos
                headers_252, campos_252 = detectar_columnas_desde_datos(table_252_data)
                table_252_preparada = preparar_datos_tabla(
                    table_252_data, headers_252, campos_252,
                    agregar_totales=False, columna_total_texto=0
                )
                contexto["table_252_1"] = table_252_preparada
                # Contar registros de la tabla
                contexto["count_rows_252_1"] = count_rows(table_252_preparada)
                
                # Agrupar por CONSECUTIVO CAÍDA y contar número de consecutivos únicos
                group_by_consecutivo = self._groupby_consecutivo_caida(table_252_data)
                num_consecutivos_unicos = len(group_by_consecutivo)
                contexto["suma_total_registros_252"] = num_consecutivos_unicos
            else:
                contexto["table_252_1"] = []
                contexto["count_rows_252_1"] = 0
                contexto["suma_total_registros_252"] = 0
        
        # Sección 2.5.3 (index[8]) - CONECTIVIDAD
        if len(index) > 8:
            content_253 = index[8].get("content", {})
            table_253_data = content_253.get("table_1", [])
            if table_253_data:
                # Detectar columnas automáticamente desde los datos
                headers_253, campos_253 = detectar_columnas_desde_datos(table_253_data)
                table_253_preparada = preparar_datos_tabla(
                    table_253_data, headers_253, campos_253,
                    agregar_totales=False, columna_total_texto=0
                )
                contexto["table_253_1"] = table_253_preparada
                # Contar registros de la tabla
                contexto["count_rows_253_1"] = count_rows(table_253_preparada)
            else:
                contexto["table_253_1"] = []
                contexto["count_rows_253_1"] = 0
        
        # Sección 2.6 (index[9]) - INFORME ACTUALIZADO DE HOJAS DE VIDA
        if len(index) > 9:
            content_26 = index[9].get("content", {})
            contexto["name_document_26"] = content_26.get("name_document", "")
        
        # Sección 2.7 (index[10]) - INFORME EJECUTIVO DEL ESTADO DEL SISTEMA
        if len(index) > 10:
            content_27 = index[10].get("content", {})
            
            # Preparar table_27_1 (ESTADO - CANTIDAD)
            table_27_1_data = content_27.get("table_1", [])
            if table_27_1_data:
                headers_27_1 = ["ESTADO", "CANTIDAD"]
                campos_27_1 = ["estado", "cantidad"]
                contexto["table_27_1"] = preparar_datos_tabla(
                    table_27_1_data, headers_27_1, campos_27_1,
                    agregar_totales=False, columna_total_texto=0
                )
            else:
                contexto["table_27_1"] = []
            
            contexto["image_27"] = procesar_imagen(template, content_27.get("image", "")) or ""
            contexto["section_27_1"] = content_27.get("section_1", "")
            contexto["section_27_2"] = content_27.get("section_2", "")
            
            # Preparar table_27_2 (RESPONSABLE - CANTIDAD)
            table_27_2_data = content_27.get("table_2", [])
            if table_27_2_data:
                headers_27_2 = ["RESPONSABLE", "CANTIDAD"]
                campos_27_2 = ["responsable", "cantidad"]
                contexto["table_27_2"] = preparar_datos_tabla(
                    table_27_2_data, headers_27_2, campos_27_2,
                    agregar_totales=False, columna_total_texto=0
                )
            else:
                contexto["table_27_2"] = []
            
            contexto["section_27_3"] = content_27.get("section_3", "")
            
            # Preparar table_27_3 (SUBSISTEMAS con múltiples columnas)
            table_27_3_data = content_27.get("table_3", [])
            if table_27_3_data:
                headers_27_3 = ["SUBSISTEMAS", "CAÍDA MASIVA", "FUERA DE SERVICIO", "OPERATIVA", "OPERATIVA CON NOVEDAD", "TOTAL"]
                campos_27_3 = ["subsistema", "caida_masiva", "fuera_de_servicio", "operativa", "operativa_con_novedad", "total"]
                contexto["table_27_3"] = preparar_datos_tabla(
                    table_27_3_data, headers_27_3, campos_27_3,
                    agregar_totales=False, columna_total_texto=0
                )
            else:
                contexto["table_27_3"] = []
            contexto["observaciones_27"] = content_27.get("observaciones", "")
            contexto["name_document_27"] = content_27.get("name_document", "")
            
            # Obtener el último día del mes en formato "30 de SEPTIEMBRE de 2025"
            if anio and mes:
                contexto["last_day_27"] = self._obtener_ultimo_dia_mes(mes, anio)
            else:
                contexto["last_day_27"] = ""

        
        return contexto
    
    def generar(self, document: Dict[str, Any], output_path: Optional[Path] = None):       
        
        template_path = config.TEMPLATES_DIR / self.template_file
        template = DocxTemplate(str(template_path) )
                
        contexto = self._construir_contexto(document , template)
        if output_path:
            # Asegurar que el output_path tenga extensión .docx
            if not str(output_path).endswith('.docx'):
                output_path = Path(str(output_path) + '.docx')

        output_path.parent.mkdir(parents=True, exist_ok=True)

        # Preparar placeholders especiales para tablas
        # El usuario puede usar {{ table_21_1_placeholder }} o [[TABLE_21_1]] directamente en el Word
        contexto_tablas = contexto.copy()
        if contexto.get("table_21_1"):
            contexto_tablas["table_21_1_placeholder"] = "[[TABLE_21_1]]"
        if contexto.get("table_21_2"):
            contexto_tablas["table_21_2_placeholder"] = "[[TABLE_21_2]]"
        if contexto.get("table_23_1"):
            contexto_tablas["table_23_1_placeholder"] = "[[TABLE_23_1]]"
        if contexto.get("table_24_1"):
            contexto_tablas["table_24_1_placeholder"] = "[[TABLE_24_1]]"
        if contexto.get("table_24_2"):
            contexto_tablas["table_24_2_placeholder"] = "[[TABLE_24_2]]"
        if contexto.get("table_25_1"):
            contexto_tablas["table_25_1_placeholder"] = "[[TABLE_25_1]]"
        if contexto.get("table_251_1"):
            contexto_tablas["table_251_1_placeholder"] = "[[TABLE_251_1]]"
        if contexto.get("table_252_1"):
            contexto_tablas["table_252_1_placeholder"] = "[[TABLE_252_1]]"
        if contexto.get("table_253_1"):
            contexto_tablas["table_253_1_placeholder"] = "[[TABLE_253_1]]"
        
        # Agregar placeholders para tablas de la sección 2.7
        if contexto.get("table_27_1"):
            contexto_tablas["table_27_1_placeholder"] = "[[TABLE_27_1]]"
        if contexto.get("table_27_2"):
            contexto_tablas["table_27_2_placeholder"] = "[[TABLE_27_2]]"
        if contexto.get("table_27_3"):
            contexto_tablas["table_27_3_placeholder"] = "[[TABLE_27_3]]"
        
        # Renderizar el template con las variables básicas
        template.render(contexto_tablas)
        
        # Guardar temporalmente para poder trabajar con el documento renderizado
        temp_path = str(output_path).replace('.docx', '_temp.docx')
        template.save(temp_path)
        
        # Abrir el documento renderizado para insertar tablas programáticamente
        doc = Document(temp_path)
        
        # Reemplazar placeholders de tablas con tablas creadas programáticamente
        if contexto.get("table_21_1"):
            reemplazar_placeholder_con_tabla(doc, "[[TABLE_21_1]]", contexto["table_21_1"])
        
        if contexto.get("table_21_2"):
            reemplazar_placeholder_con_tabla(doc, "[[TABLE_21_2]]", contexto["table_21_2"])
        
        if contexto.get("table_23_1"):
            reemplazar_placeholder_con_tabla(doc, "[[TABLE_23_1]]", contexto["table_23_1"])
        
        if contexto.get("table_24_1"):
            reemplazar_placeholder_con_tabla(doc, "[[TABLE_24_1]]", contexto["table_24_1"])
        
        if contexto.get("table_24_2"):
            reemplazar_placeholder_con_tabla(doc, "[[TABLE_24_2]]", contexto["table_24_2"])
        
        if contexto.get("table_25_1"):
            reemplazar_placeholder_con_tabla(doc, "[[TABLE_25_1]]", contexto["table_25_1"])
        
        if contexto.get("table_251_1"):
            reemplazar_placeholder_con_tabla(doc, "[[TABLE_251_1]]", contexto["table_251_1"])
        
        if contexto.get("table_252_1"):
            reemplazar_placeholder_con_tabla(doc, "[[TABLE_252_1]]", contexto["table_252_1"])
        
        if contexto.get("table_253_1"):
            reemplazar_placeholder_con_tabla(doc, "[[TABLE_253_1]]", contexto["table_253_1"])
        
        # Reemplazar tablas de la sección 2.7 (Estado del Sistema)
        if contexto.get("table_27_1"):
            reemplazar_placeholder_con_tabla(doc, "[[TABLE_27_1]]", contexto["table_27_1"])
        
        if contexto.get("table_27_2"):
            reemplazar_placeholder_con_tabla(doc, "[[TABLE_27_2]]", contexto["table_27_2"])
        
        if contexto.get("table_27_3"):
            reemplazar_placeholder_con_tabla(doc, "[[TABLE_27_3]]", contexto["table_27_3"])
        
        # Guardar el documento final
        doc.save(str(output_path))
        
        # Eliminar archivo temporal
        try:
            Path(temp_path).unlink()
        except:
            pass

        print(f"[OK] Documento generado en: {output_path}")
        
        return template
